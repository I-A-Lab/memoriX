"""Generate a professional DOCX from the memoriX A/B comparison report and CSV data."""

import csv
import json
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------
SCRIPT_DIR = Path(__file__).resolve().parent
REPORT_MD = SCRIPT_DIR / "ab-comparison" / "ab-comparison-report.md"
CSV_FILE = SCRIPT_DIR / "ab-comparison" / "ab-summary.csv"
RAW_JSON = SCRIPT_DIR / "memory-benchmark-results" / "raw_results.json"
OUTPUT_DOCX = SCRIPT_DIR / "ab-comparison" / "ab-comparison-report.docx"

# ---------------------------------------------------------------------------
# Styling constants
# ---------------------------------------------------------------------------
HEADER_BG = "1E293B"   # dark slate
HEADER_FG = RGBColor(0xFF, 0xFF, 0xFF)
ALT_ROW_BG = "F1F5F9"  # light slate
TITLE_COLOR = RGBColor(0x0F, 0x17, 0x2A)
ACCENT = RGBColor(0x22, 0xD3, 0xEE)  # cyan accent
GREEN = RGBColor(0x16, 0xA3, 0x4A)
RED = RGBColor(0xDC, 0x26, 0x26)
GRAY = RGBColor(0x64, 0x74, 0x8B)

# Family ID -> (Family Name, What It Tests)
FAMILY_EXPLANATIONS = {
    "f01": ("Exact Key Recall", "Retrieval of a stored memory by its exact key. Tests precision and recall for direct lookups."),
    "f02": ("Semantic Retrieval", "Retrieval using semantic (meaning-based) similarity rather than exact key matches. Harder fuzzy matching."),
    "f03": ("Distractor Robustness", "Retrieval accuracy as the number of distractor (irrelevant) memories increases. Measures robustness against confounding information."),
    "f04": ("Multi-Hop Reasoning", "Answering questions that require chaining multiple pieces of information across several stored memories."),
    "f05": ("Temporal Decay", "Retrieval quality following an exponential decay model over time. Measures handling of naturally fading memories."),
    "f06": ("Contextual Disambiguation", "Disambiguation of vague or ambiguous queries against stored memories. Context-dependent interpretation."),
    "f07": ("Adversarial Injection", "Detection and correct handling of contradictory or adversarially injected information in stored facts."),
    "f08": ("Cross-Session Leak", "Isolation of memories between sessions. Detects leakage of information from one session into another."),
    "f09": ("Privacy Isolation", "Correct handling of sensitive or restricted information. High precision for safety-critical responses."),
    "f10": ("Consolidation Quality", "Quality of memory consolidation - ability to aggregate and synthesize scattered facts into a coherent answer."),
    "f11": ("Forgetting Curve", "How retrieval quality changes as memories age (Ebbinghaus-style forgetting curve)."),
    "f12": ("Capacity Management", "Performance as the memory store fills to capacity. Graceful degradation under pressure."),
    "f13": ("Search Precision", "Precision of search results - retrieving relevant memories without noise."),
    "f14": ("Search Recall", "Recall of search - finding all relevant stored memories for a query."),
    "f15": ("Ranking Relevance", "Quality of result ranking - most relevant memories appear first."),
    "f16": ("Query Expansion", "Ability to expand queries with synonyms or related terms to improve retrieval."),
    "f17": ("Partial Match", "Retrieval when queries only partially match stored memories. Incomplete or approximate lookups."),
    "f18": ("Noise Tolerance", "Retrieval accuracy with systematic noise injected into the memory store."),
    "f19": ("Schema Evolution", "Handling of schema changes over time. Adaptation to evolving memory structures."),
    "f20": ("Version Awareness", "Awareness of memory versions. Correct handling of outdated vs current memories."),
    "f21": ("Concurrent Access", "Performance under concurrent access patterns. Data integrity during simultaneous read/write."),
    "f22": ("Migration Safety", "Safe migration of memories between storage systems without data loss or corruption."),
    "f23": ("API Conformance", "Conformance to the public API contract. Correct responses to API calls."),
    "f24": ("SDK Compatibility", "Compatibility with SDK usage patterns. Correct integration with client libraries."),
    "f25": ("Performance Baseline", "Baseline performance of the memory system under normal conditions."),
    "f26": ("Latency Budget", "Retrieval latency within budget as memory count grows. Real-time responsiveness."),
    "f27": ("Memory Footprint", "Memory usage efficiency. Retrieval quality with minimal storage overhead."),
    "f28": ("Streaming Integrity", "Correct handling of streaming operations. No data corruption during incremental processing."),
    "f29": ("End-to-End Smoke", "Complete pipeline from ingestion to retrieval in an integrated scenario."),
    "f30": ("Reproducibility", "Consistency of results across repeated runs with the same inputs."),
    "f31": ("Determinism Check", "Determinism of outputs - same input always produces same output."),
    "f32": ("Golden Output", "Matching of outputs to golden/reference outputs defined by the task spec."),
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def set_cell_bg(cell, color_hex: str):
    """Set cell background colour via shading XML."""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_text(cell, text, bold=False, color=None, size=9, alignment=None):
    """Write styled text into a cell."""
    cell.text = ""
    p = cell.paragraphs[0]
    if alignment is not None:
        p.alignment = alignment
    run = p.add_run(str(text))
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if bold:
        run.bold = True
    if color:
        run.font.color.rgb = color
    # Reduce cell padding
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)


def style_header_row(row, num_cols):
    """Dark background + white bold text for header row."""
    for i in range(num_cols):
        set_cell_bg(row.cells[i], HEADER_BG)
        set_cell_text(row.cells[i], row.cells[i].text, bold=True, color=HEADER_FG, size=9)


def add_styled_table(doc, headers, rows, col_widths=None):
    """Create a table with header styling and alternating row colours."""
    num_cols = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=num_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    # Header row
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        set_cell_text(hdr.cells[i], h, bold=True, color=HEADER_FG, size=9)
        set_cell_bg(hdr.cells[i], HEADER_BG)

    # Data rows
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        for c_idx, val in enumerate(row_data):
            set_cell_text(row.cells[c_idx], val, size=9)
            if r_idx % 2 == 1:
                set_cell_bg(row.cells[c_idx], ALT_ROW_BG)

    # Column widths
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)

    return table


def add_heading(doc, text, level=1):
    """Add a heading with Calibri."""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = "Calibri"
    return heading


def add_body(doc, text, bold=False, color=None, size=10):
    """Add a normal paragraph."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    if bold:
        run.bold = True
    if color:
        run.font.color.rgb = color
    return p


# ---------------------------------------------------------------------------
# CSV parsing
# ---------------------------------------------------------------------------
def read_csv():
    """Return headers + sorted rows from ab-summary.csv."""
    with open(CSV_FILE, newline="", encoding="utf-8") as f:
        reader = csv.reader(f)
        headers_raw = next(reader)
        data = []
        for row in reader:
            data.append(row)
    # Sort by delta_pp descending (column index 7, as string -> float)
    data.sort(key=lambda r: float(r[7]), reverse=True)
    return headers_raw, data


def load_raw_results():
    """Return the list of run dicts from raw_results.json."""
    with open(RAW_JSON, encoding="utf-8") as f:
        return json.load(f)


def median(values):
    """Median of a list of numbers."""
    s = sorted(values)
    n = len(s)
    if n == 0:
        return 0.0
    mid = n // 2
    if n % 2 == 1:
        return s[mid]
    return (s[mid - 1] + s[mid]) / 2


def compute_overall_stats(runs):
    """Return {mode: (passed, total)} for the two benchmark modes."""
    stats = {}
    for mode in ("no_memory", "memorix_core"):
        mode_runs = [r for r in runs if r["mode"] == mode]
        stats[mode] = (sum(1 for r in mode_runs if r["passed"]), len(mode_runs))
    return stats


def read_csv_summary():
    """Return ab-summary.csv rows as dicts with numeric fields converted."""
    with open(CSV_FILE, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        rows = []
        for row in reader:
            rows.append({
                "family": row["family"],
                "no_memory_passed": int(row["no_memory_passed"]),
                "no_memory_total": int(row["no_memory_total"]),
                "no_memory_rate": float(row["no_memory_rate"]),
                "memorix_core_passed": int(row["memorix_core_passed"]),
                "memorix_core_total": int(row["memorix_core_total"]),
                "memorix_core_rate": float(row["memorix_core_rate"]),
                "delta_pp": float(row["delta_pp"]),
                "median_lat_no_memory_ms": float(row["median_lat_no_memory_ms"]),
                "median_lat_memorix_core_ms": float(row["median_lat_memorix_core_ms"]),
            })
    return rows


# ---------------------------------------------------------------------------
# Markdown parsing for key-findings
# ---------------------------------------------------------------------------
def parse_key_findings(md_text):
    """Extract the three key-findings subsections."""
    sections = {}
    pattern = r"### (.*?)\n\n(.*?)(?=\n### |\n## |\Z)"
    for match in re.finditer(pattern, md_text, re.DOTALL):
        title = match.group(1).strip()
        body = match.group(2).strip()
        items = [
            line.strip().lstrip("- ").strip()
            for line in body.split("\n")
            if line.strip().startswith("-")
        ]
        sections[title] = items
    return sections


# ---------------------------------------------------------------------------
# Build DOCX
# ---------------------------------------------------------------------------
def build_document():
    doc = Document()

    # -- Page setup --
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    # -- Default font --
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(10)

    # ====================================================================
    # TITLE
    # ====================================================================
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_after = Pt(4)
    run = title_p.add_run("memoriX Benchmark")
    run.bold = True
    run.font.size = Pt(26)
    run.font.name = "Calibri"
    run.font.color.rgb = TITLE_COLOR

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_p.paragraph_format.space_after = Pt(16)
    run = subtitle_p.add_run("A/B Comparison Report")
    run.bold = True
    run.font.size = Pt(18)
    run.font.name = "Calibri"
    run.font.color.rgb = GRAY

    # Thin divider
    div = doc.add_paragraph()
    div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    div.paragraph_format.space_before = Pt(0)
    div.paragraph_format.space_after = Pt(12)
    run = div.add_run("─" * 60)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)

    # ====================================================================
    # CAMPAIGN INFO
    # ====================================================================
    add_heading(doc, "Campaign Information", level=2)
    runs = load_raw_results()
    total_runs = len(runs)
    ab_pairs = total_runs // 2
    n_seeds = len({r["seed"] for r in runs})
    n_families = len({r["family"] for r in runs})
    campaign_data = [
        ("Campaign", "memory-benchmark"),
        ("Total runs", f"{total_runs:,} ({n_seeds} seeds x {n_families} families x 2 modes)"),
        ("A/B pairs", f"{ab_pairs:,}"),
        ("Model", "qwen2.5:3b (Ollama, local)"),
        ("Modes", "memorix_core (with-memory) vs no_memory (baseline)"),
        ("Seeds per family", f"{n_seeds}"),
        ("Families", f"{n_families}"),
    ]
    table = doc.add_table(rows=len(campaign_data), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, (label, value) in enumerate(campaign_data):
        set_cell_text(table.rows[i].cells[0], label, bold=True, size=9)
        set_cell_bg(table.rows[i].cells[0], "F8FAFC")
        set_cell_text(table.rows[i].cells[1], value, size=9)
        table.rows[i].cells[0].width = Cm(4.5)
        table.rows[i].cells[1].width = Cm(12)
        if i % 2 == 1:
            set_cell_bg(table.rows[i].cells[1], ALT_ROW_BG)

    doc.add_paragraph()  # spacer

    # ====================================================================
    # EXECUTIVE SUMMARY
    # ====================================================================
    add_heading(doc, "Executive Summary", level=2)

    summary_headers = ["Metric", "No-Memory", "With-Memory (memorix_core)", "Delta"]
    csv_summary = read_csv_summary()
    overall = compute_overall_stats(runs)
    nm_passed, nm_total = overall["no_memory"]
    mc_passed, mc_total = overall["memorix_core"]
    nm_rate = 100.0 * nm_passed / nm_total
    mc_rate = 100.0 * mc_passed / mc_total
    med_lat_no = median([r["median_lat_no_memory_ms"] for r in csv_summary])
    med_lat_mc = median([r["median_lat_memorix_core_ms"] for r in csv_summary])
    fam_won = sum(1 for r in csv_summary if r["delta_pp"] > 0)
    fam_lost = sum(1 for r in csv_summary if r["delta_pp"] < 0)
    fam_tied = sum(1 for r in csv_summary if r["delta_pp"] == 0)
    summary_rows = [
        ["Overall pass rate", f"{nm_rate:.1f}% ({nm_passed}/{nm_total})", f"{mc_rate:.1f}% ({mc_passed}/{mc_total})", f"{mc_rate - nm_rate:+.1f}pp"],
        ["Median family latency", f"{med_lat_no:.0f} ms", f"{med_lat_mc:.0f} ms", f"{med_lat_mc - med_lat_no:+.0f} ms"],
        ["Families evaluated", f"{n_families}", f"{n_families}", "--"],
        ["Families won (delta > 0)", "--", f"{fam_won}", "--"],
        ["Families lost (delta < 0)", "--", f"{fam_lost}", "--"],
        ["Families tied (delta = 0)", "--", f"{fam_tied}", "--"],
    ]
    add_styled_table(doc, summary_headers, summary_rows, col_widths=[4.5, 4, 4.5, 2.5])

    doc.add_paragraph()  # spacer

    # ====================================================================
    # PER-FAMILY A/B COMPARISON TABLE
    # ====================================================================
    add_heading(doc, "Per-Family A/B Comparison", level=2)
    add_body(doc, "All 32 families sorted by delta (descending). Families with positive delta show improvement with memory.", size=9, color=GRAY)

    csv_headers, csv_rows = read_csv()

    # Build display rows from CSV
    family_headers = [
        "Family",
        "No-Memory Rate",
        "With-Memory Rate",
        "Delta (pp)",
        "Median Lat No-Mem (ms)",
        "Median Lat With-Mem (ms)",
    ]
    family_display_rows = []
    for row in csv_rows:
        family_name = row[0]
        no_mem_rate = f"{float(row[3]):.1f}% ({row[1]}/{row[2]})"
        with_mem_rate = f"{float(row[6]):.1f}% ({row[4]}/{row[5]})"
        delta_val = float(row[7])
        delta_str = f"{delta_val:+.1f}"
        lat_no = f"{float(row[8]):.0f}"
        lat_with = f"{float(row[9]):.0f}"
        family_display_rows.append([family_name, no_mem_rate, with_mem_rate, delta_str, lat_no, lat_with])

    add_styled_table(doc, family_headers, family_display_rows, col_widths=[4.2, 2.8, 2.8, 1.8, 2.2, 2.2])

    # Color-code delta cells after table creation
    # (We need to re-iterate the table to find delta cells)
    table = doc.tables[-1]  # last table added
    for r_idx in range(1, len(table.rows)):
        delta_cell = table.rows[r_idx].cells[3]
        try:
            delta_val = float(delta_cell.text)
        except ValueError:
            continue
        if delta_val > 0:
            set_cell_text(delta_cell, delta_cell.text, bold=True, color=GREEN, size=9)
        elif delta_val < 0:
            set_cell_text(delta_cell, delta_cell.text, bold=True, color=RED, size=9)

    doc.add_paragraph()  # spacer

    # ====================================================================
    # KEY FINDINGS
    # ====================================================================
    add_heading(doc, "Key Findings", level=2)

    with open(REPORT_MD, encoding="utf-8") as f:
        md_text = f.read()

    findings = parse_key_findings(md_text)

    # --- Strong improvement ---
    strong_key = [k for k in findings if "Strong" in k]
    if strong_key:
        add_heading(doc, f"Strong Improvement ({strong_key[0].split('(')[1].rstrip(')') if '(' in strong_key[0] else ''})", level=3)
        for item in findings[strong_key[0]]:
            p = doc.add_paragraph(style="List Bullet")
            # Parse bold family name and delta
            m = re.match(r"\*\*(.+?)\*\*: (.+)", item)
            if m:
                family_name = m.group(1)
                rest = m.group(2)
                run1 = p.add_run(family_name)
                run1.bold = True
                run1.font.name = "Calibri"
                run1.font.size = Pt(10)
                run2 = p.add_run(f"  {rest}")
                run2.font.name = "Calibri"
                run2.font.size = Pt(10)
                run2.font.color.rgb = GREEN
            else:
                run = p.add_run(item)
                run.font.name = "Calibri"
                run.font.size = Pt(10)

    # --- Moderate improvement ---
    moderate_key = [k for k in findings if "Moderate" in k]
    if moderate_key:
        add_heading(doc, f"Moderate Improvement ({moderate_key[0].split('(')[1].rstrip(')') if '(' in moderate_key[0] else ''})", level=3)
        for item in findings[moderate_key[0]]:
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(item)
            run.font.name = "Calibri"
            run.font.size = Pt(10)

    # --- No improvement / regression ---
    regress_key = [k for k in findings if "No Significant" in k or "Regression" in k]
    if regress_key:
        add_heading(doc, f"No Significant Improvement / Regression ({regress_key[0].split('(')[1].rstrip(')') if '(' in regress_key[0] else ''})", level=3)
        for item in findings[regress_key[0]]:
            p = doc.add_paragraph(style="List Bullet")
            m = re.match(r"\*\*(.+?)\*\*: (.+)", item)
            if m:
                family_name = m.group(1)
                rest = m.group(2)
                is_regression = "REGRESSION" in rest
                run1 = p.add_run(family_name)
                run1.bold = True
                run1.font.name = "Calibri"
                run1.font.size = Pt(10)
                run2 = p.add_run(f"  {rest}")
                run2.font.name = "Calibri"
                run2.font.size = Pt(10)
                run2.font.color.rgb = RED if is_regression else GRAY
            else:
                run = p.add_run(item)
                run.font.name = "Calibri"
                run.font.size = Pt(10)

    doc.add_paragraph()  # spacer

    # ====================================================================
    # FAMILY EXPLANATIONS
    # ====================================================================
    add_heading(doc, "Family Explanations", level=2)
    add_body(doc, "Description of what each benchmark family evaluates and what it tests.", size=9, color=GRAY)

    family_expl_headers = ["Family ID", "Family Name", "What It Tests"]
    family_expl_rows = []
    for family_key in sorted({r["family"] for r in runs}):
        family_id = family_key.split("_", 1)[0]
        family_name, description = FAMILY_EXPLANATIONS.get(family_id, (family_key, "No description available."))
        family_expl_rows.append([family_id, family_name, description])
    add_styled_table(doc, family_expl_headers, family_expl_rows, col_widths=[2.0, 4.2, 10.4])

    doc.add_paragraph()  # spacer

    # ====================================================================
    # METHODOLOGY
    # ====================================================================
    add_heading(doc, "Methodology", level=2)

    methodology_items = [
        ("Approach", "Prompt injection -- the LLM receives the full task prompt along with retrieved memory context (memorix_core) or without (no_memory)."),
        ("Model", "qwen2.5:3b served locally via Ollama."),
        ("Design", f"{n_seeds} seeds x {n_families} families x 2 modes = {total_runs} total runs."),
        ("Pass criteria", "All test cases for a given seed/family must pass for the run to be marked passed."),
        ("Latency", "Wall-clock LLM inference time per run (llm_ms)."),
    ]

    for label, description in methodology_items:
        p = doc.add_paragraph()
        run_label = p.add_run(f"{label}: ")
        run_label.bold = True
        run_label.font.name = "Calibri"
        run_label.font.size = Pt(10)
        run_desc = p.add_run(description)
        run_desc.font.name = "Calibri"
        run_desc.font.size = Pt(10)

    doc.add_paragraph()  # spacer

    # ====================================================================
    # RAW DATA REFERENCES
    # ====================================================================
    add_heading(doc, "Raw Data", level=2)
    p = doc.add_paragraph()
    run = p.add_run("Source file: ")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run = p.add_run(r"C:\GitHub\memoriX\.benchmarks_v2\memory_run\memory-benchmark-results\raw_results.json")
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY

    p = doc.add_paragraph()
    run = p.add_run("CSV summary: ")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run = p.add_run(str(CSV_FILE))
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY

    # Footer divider
    doc.add_paragraph()
    div = doc.add_paragraph()
    div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = div.add_run("─" * 60)
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)

    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_p.add_run("Report generated by generate_ab_docx.py")
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY
    run.italic = True

    # ====================================================================
    # SAVE
    # ====================================================================
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUTPUT_DOCX))
    print(f"DOCX written to: {OUTPUT_DOCX}")


# ---------------------------------------------------------------------------
if __name__ == "__main__":
    build_document()
